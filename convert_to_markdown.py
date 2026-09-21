import os
from pathlib import Path
import subprocess
import warnings
import xml.etree.ElementTree as ET
import zipfile

# Silence openpyxl warnings
warnings.filterwarnings("ignore")

import docx
import pandas as pd
from pptx import Presentation
from pypdf import PdfReader

import pandas as pd

# --- EXTRACTION HELPERS ---


def extract_txt(file_path):
    return file_path.read_text(encoding="utf-8", errors="ignore")


def extract_docx(file_path):
    doc = docx.Document(file_path)
    md_output = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            if p.style.name.startswith("Heading"):
                md_output.append(f"### {text}\n")
            else:
                md_output.append(f"{text}\n")

    for table in doc.tables:
        table_rows = []
        for i, row in enumerate(table.rows):
            cells = [
                cell.text.strip().replace("\n", " ") for cell in row.cells
            ]
            table_rows.append("| " + " | ".join(cells) + " |")
            if i == 0:
                table_rows.append("| " + " | ".join(["---"] * len(cells)) + " |")
        if table_rows:
            md_output.append("\n" + "\n".join(table_rows) + "\n")

    return "\n".join(md_output)


def extract_doc(file_path):
    try:
        temp_dir = file_path.parent / "_temp_doc_conv"
        temp_dir.mkdir(exist_ok=True)
        subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "docx",
                str(file_path),
                "--outdir",
                str(temp_dir),
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            check=True,
        )
        converted_file = temp_dir / (file_path.stem + ".docx")
        if converted_file.exists():
            text = extract_docx(converted_file)
            converted_file.unlink()
            return text
    except Exception:
        pass
    return f"[Warning: Could not parse legacy .doc file {file_path.name}. Convert to .docx manually.]"


def extract_excel(file_path, max_cols_for_table=15):
    """Extracts Excel spreadsheets cleanly, preventing giant column explosions."""
    ext = file_path.suffix.lower()

    if ext == ".csv":
        sheets = {"Sheet1": pd.read_csv(file_path)}
    else:
        sheets = pd.read_excel(file_path, sheet_name=None, engine="openpyxl")

    md_output = []

    for sheet_name, df in sheets.items():
        # 1. Drop completely empty rows and columns
        df = df.dropna(how="all", axis=0).dropna(how="all", axis=1)

        if df.empty:
            continue

        # 2. Filter out 'Unnamed' columns that are completely empty
        cols_to_keep = [
            c
            for c in df.columns
            if not (str(c).startswith("Unnamed:") and df[c].isna().all())
        ]
        df = df[cols_to_keep]

        if df.empty:
            continue

        md_output.append(f"## Sheet: {sheet_name}\n")

        # 3. Handle extra-wide tables (e.g., > 15 columns) safely
        if len(df.columns) > max_cols_for_table:
            # Convert wide rows into readable key-value blocks instead of a 100+ column Markdown line
            records = []
            for idx, row in df.iterrows():
                row_items = [
                    f"  - **{col}**: {val}"
                    for col, val in row.items()
                    if pd.notna(val) and str(val).strip() != ""
                ]
                if row_items:
                    records.append(
                        f"### Row {idx + 1}\n" + "\n".join(row_items)
                    )
            md_output.append("\n\n".join(records) + "\n")
        else:
            # Standard size tables stay as Markdown tables
            md_table = df.fillna("").to_markdown(index=False)
            md_output.append(md_table + "\n")

    return "\n".join(md_output)


def extract_potx_fallback(file_path):
    """Fallback XML parser for PowerPoint .potx templates."""
    md_output = []
    try:
        with zipfile.ZipFile(file_path, "r") as z:
            slide_files = sorted([
                f
                for f in z.namelist()
                if f.startswith("ppt/slides/slide") and f.endswith(".xml")
            ])
            for idx, s_file in enumerate(slide_files, start=1):
                tree = ET.fromstring(z.read(s_file))
                texts = [
                    elem.text
                    for elem in tree.iter()
                    if elem.tag.endswith("}t") and elem.text
                ]
                if texts:
                    md_output.append(
                        f"### Slide {idx}\n" + "\n".join(texts) + "\n"
                    )
    except Exception as e:
        return f"[Warning: Could not parse .potx file: {e}]"
    return (
        "\n".join(md_output)
        if md_output
        else "[Template file contains no text slides]"
    )


def extract_pptx(file_path):
    try:
        prs = Presentation(file_path)
        md_output = []
        for idx, slide in enumerate(prs.slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            if slide_text:
                md_output.append(
                    f"### Slide {idx}\n" + "\n".join(slide_text) + "\n"
                )
        return "\n".join(md_output)
    except Exception:
        # Fallback to direct XML parsing if python-pptx rejects .potx header
        return extract_potx_fallback(file_path)


def extract_pdf(file_path):
    reader = PdfReader(file_path)
    md_output = []
    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text()
        if text and text.strip():
            md_output.append(f"### Page {idx}\n{text.strip()}\n")
    return "\n".join(md_output)


# --- MAIN LOGIC ---


def process_all_extensions(
    source_dir="resources",
    output_dir="./merged_md_output",
    max_files_per_batch=100,
):
    source_path = Path(source_dir)
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    ignored_extensions = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".svg"}

    extractors = {
        ".txt": extract_txt,
        ".docx": extract_docx,
        ".doc": extract_doc,
        ".xlsx": extract_excel,
        ".xlsm": extract_excel,
        ".csv": extract_excel,
        ".pptx": extract_pptx,
        ".potx": extract_pptx,
        ".pdf": extract_pdf,
    }

    files_by_ext = {}
    for f in source_path.rglob("*"):
        if f.is_file():
            ext = f.suffix.lower()
            if ext in ignored_extensions or ext not in extractors:
                continue
            files_by_ext.setdefault(ext, []).append(f)

    print(
        f"Found {sum(len(v) for v in files_by_ext.values())} processable files across {len(files_by_ext)} extensions.\n"
    )

    for ext, file_list in files_by_ext.items():
        clean_ext_name = ext.replace(".", "")
        extractor = extractors[ext]
        print(f"Processing {len(file_list)} files for extension: '{ext}'...")

        batch_num = 1
        files_in_batch = 0
        current_out_file = None

        for file_path in file_list:
            if files_in_batch == 0:
                out_name = (
                    output_path / f"merged_{clean_ext_name}_part{batch_num:02d}.md"
                )
                current_out_file = open(out_name, "w", encoding="utf-8")

            try:
                relative_path = file_path.relative_to(source_path)

                current_out_file.write(f"\n\n{'='*80}\n")
                current_out_file.write(f"# FILE: {relative_path}\n")
                current_out_file.write(f"{'='*80}\n\n")

                content = extractor(file_path)
                current_out_file.write(content + "\n\n")

                files_in_batch += 1

                if files_in_batch >= max_files_per_batch:
                    current_out_file.close()
                    batch_num += 1
                    files_in_batch = 0

            except Exception as e:
                print(f"  [Skipped] {file_path.name}: {e}")

        if current_out_file and not current_out_file.closed:
            current_out_file.close()

    print("\nAll done! Converted files are saved in:", output_path.resolve())


if __name__ == "__main__":
    process_all_extensions(
        source_dir="resources", output_dir="./merged_md_output"
    )

